import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.errors import AppError
from app.safety import text_looks_like_injection

PAGE_SPLIT = re.compile(r"\[PAGE\s+(\d+)\]")
SECTION_RE = re.compile(r"\[SECTION\s+([a-z0-9][a-z0-9-]*)\]\s*(.*)")
TABLE_LINE = re.compile(r"^\s*\|.+\|\s*$")


@dataclass
class SectionBlock:
    section_id: str
    title: str
    text: str
    has_table: bool = False
    injection_suspected: bool = False


@dataclass
class ParsedPage:
    page: int
    sections: list[SectionBlock] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        reader = PdfReader(BytesIO(data))
        if len(reader.pages) > 500:
            raise AppError("file_too_large", "PDFs are limited to 500 pages.", 400)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == "docx":
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext in {"txt", "md", "markdown"}:
        return data.decode("utf-8", errors="replace")
    raise AppError("unsupported_type", "Upload a PDF, DOCX, TXT, or Markdown file.", 400)


def _tables_from_lines(lines: list[str]) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    for line in lines:
        if TABLE_LINE.match(line) and "---" not in line:
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            current.append(cells)
        elif current:
            if len(current) >= 2:
                tables.append(current)
            current = []
    if len(current) >= 2:
        tables.append(current)
    return tables


def parse_marked_text(text: str) -> list[ParsedPage]:
    if len(text) > 2_000_000:
        raise AppError("file_too_large", "Extracted text exceeds the parser limit.", 400)
    parts = PAGE_SPLIT.split(text)
    pages: list[ParsedPage] = []
    if len(parts) == 1:
        pages.append(_page_from_body(1, parts[0]))
        return [page for page in pages if page.sections]
    preamble = parts[0].strip()
    index = 1
    while index < len(parts) - 1:
        page_no = int(parts[index])
        body = parts[index + 1]
        if index == 1 and preamble:
            body = preamble + "\n" + body
        pages.append(_page_from_body(page_no, body))
        index += 2
    return [page for page in pages if page.sections]


def _page_from_body(page_no: int, body: str) -> ParsedPage:
    lines = body.splitlines()
    tables = _tables_from_lines(lines)
    sections: list[SectionBlock] = []
    current_id = f"page-{page_no}"
    current_title = f"Page {page_no}"
    buffer: list[str] = []
    seen_section = False

    def flush() -> None:
        text = "\n".join(buffer).strip()
        if not text:
            return
        sections.append(
            SectionBlock(
                section_id=current_id,
                title=current_title.strip() or current_id,
                text=text,
                has_table=any(TABLE_LINE.match(line) for line in text.splitlines()),
                injection_suspected=text_looks_like_injection(text),
            )
        )

    for line in lines:
        match = SECTION_RE.match(line.strip())
        if match:
            if seen_section:
                flush()
            buffer = []
            seen_section = True
            current_id = match.group(1)
            current_title = match.group(2).strip() or match.group(1)
            continue
        buffer.append(line)
    flush()
    if not sections and body.strip():
        sections.append(
            SectionBlock(
                section_id=f"page-{page_no}",
                title=f"Page {page_no}",
                text=body.strip(),
                has_table=bool(tables),
                injection_suspected=text_looks_like_injection(body),
            )
        )
    return ParsedPage(page=page_no, sections=sections, tables=tables)


def parse_file(filename: str, data: bytes) -> list[ParsedPage]:
    return parse_marked_text(extract_text(filename, data))
