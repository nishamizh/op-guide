import uuid
from datetime import UTC, date, datetime
from io import BytesIO

from docx import Document as DocxDocument
from fpdf import FPDF
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ingestion.pipeline import process_job
from app.models import AppMeta, Document, DocumentVersion, IngestionJob
from app.seed.content import manuals
from app.seed.ids import SEED_VERSION, document_id, version_id
from app.storage import get_storage, sha256_bytes


def render_manual(manual: dict) -> tuple[str, str, bytes]:
    text = _marked_text(manual)
    kind = manual["format"]
    if kind == "pdf":
        return f"{manual['key']}.pdf", "application/pdf", _pdf(text)
    if kind == "docx":
        return (
            f"{manual['key']}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx(text),
        )
    if kind == "txt":
        return f"{manual['key']}.txt", "text/plain", text.encode()
    return f"{manual['key']}.md", "text/markdown", text.encode()


def _marked_text(manual: dict) -> str:
    parts = [
        f"DEMO {manual['title']}",
        f"Revision {manual['revision']}. Effective {manual['effective_date']}.",
    ]
    for page in manual["pages"]:
        parts.append(f"[PAGE {page['page']}]")
        for section in page["sections"]:
            parts.append(f"[SECTION {section['id']}] {section['title']}")
            parts.append(section["body"])
    return "\n".join(parts) + "\n"


def _pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    safe = text.encode("latin-1", errors="replace").decode("latin-1")
    pdf.multi_cell(0, 5, safe)
    return bytes(pdf.output())


def _docx(text: str) -> bytes:
    document = DocxDocument()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def seed_demo(db: Session, force: bool = False) -> dict:
    meta = db.get(AppMeta, 1)
    if meta is None:
        meta = AppMeta(id=1, seed_version="", index_revision=0, embedding_dimensions=1536)
        db.add(meta)
        db.flush()
    if meta.seed_version == SEED_VERSION and not force:
        count = db.execute(select(Document).where(Document.is_demo.is_(True))).all()
        if count:
            return {"seeded": False, "reason": "already-current", "documents": len(count)}
    for manual in manuals():
        doc_id = document_id(manual["key"])
        existing = db.get(Document, doc_id)
        if existing is not None:
            db.delete(existing)
            db.flush()
    storage = get_storage()
    created = []
    for manual in manuals():
        filename, mime, data = render_manual(manual)
        doc_id = document_id(manual["key"])
        ver_id = version_id(manual["key"])
        key = f"demo/{manual['key']}/{filename}"
        storage.put(key, data, mime)
        document = Document(
            id=doc_id,
            tenant_id="demo",
            user_id="seed",
            title=manual["title"],
            manufacturer=manual["manufacturer"],
            machine_model=manual["model"],
            equipment_type=manual["equipment_type"],
            plant=manual["plant"],
            document_revision=manual["revision"],
            effective_date=date.fromisoformat(manual["effective_date"]),
            filename=filename,
            mime_type=mime,
            storage_key=key,
            size_bytes=len(data),
            status="parse",
            is_demo=True,
            current_version_id=None,
            created_at=datetime.now(UTC),
            updated_at=datetime.now(UTC),
        )
        db.add(document)
        db.flush()
        version = DocumentVersion(
            id=ver_id,
            document_id=doc_id,
            revision=manual["revision"],
            effective_date=document.effective_date,
            storage_key=key,
            checksum=sha256_bytes(data),
            is_current=True,
        )
        db.add(version)
        db.flush()
        document.current_version_id = ver_id
        job = IngestionJob(
            id=uuid.uuid4(),
            tenant_id="demo",
            user_id="seed",
            document_id=doc_id,
            status="queued",
            stage="upload",
            progress=5,
        )
        db.add(job)
        db.flush()
        process_job(db, job.id)
        created.append({"id": str(doc_id), "title": manual["title"], "status": document.status})
    meta.seed_version = SEED_VERSION
    db.flush()
    return {"seeded": True, "documents": created, "seed_version": SEED_VERSION}
